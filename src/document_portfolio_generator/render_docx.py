from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .models import Profile


HEADING_BLUE = RGBColor(46, 116, 181)
HEADING_DARK_BLUE = RGBColor(31, 77, 120)
MUTED_GOLD = RGBColor(122, 90, 0)


def write_docx(profile: Profile, destination: Path) -> None:
    document = Document()
    _configure_compact_reference_styles(document)
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    document.core_properties.title = f"{profile.name} Portfolio"
    document.core_properties.author = "Document Portfolio Generator"

    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.paragraph_format.space_after = Pt(4)
    name_run = heading.add_run(profile.name)
    name_run.bold = True
    name_run.font.size = Pt(24)
    name_run.font.color.rgb = HEADING_DARK_BLUE
    _set_run_font(name_run)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(6)
    subtitle_run = subtitle.add_run(profile.headline)
    subtitle_run.font.size = Pt(12)
    subtitle_run.font.color.rgb = MUTED_GOLD
    _set_run_font(subtitle_run)
    if profile.location:
        subtitle.add_run(f" · {profile.location}")

    contact = document.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact.paragraph_format.space_after = Pt(10)
    contact.add_run(" · ".join(f"{item.label}: {item.value}" for item in profile.contact))

    _section_heading(document, "Profile")
    document.add_paragraph(profile.summary)

    _section_heading(document, "Skills")
    for section_name, values in profile.skills.items():
        paragraph = document.add_paragraph()
        paragraph.add_run(f"{section_name}: ").bold = True
        paragraph.add_run(", ".join(values))

    _section_heading(document, "Experience")
    for item in profile.experience:
        paragraph = document.add_paragraph()
        paragraph.add_run(item.role).bold = True
        paragraph.add_run(f" — {item.organization}")
        period = document.add_paragraph(item.period)
        period.runs[0].italic = True
        for highlight in item.highlights:
            bullet = document.add_paragraph(highlight, style="List Bullet")
            bullet.paragraph_format.left_indent = Inches(0.375)
            bullet.paragraph_format.first_line_indent = Inches(-0.188)
            bullet.paragraph_format.space_after = Pt(2)
            bullet.paragraph_format.line_spacing = 1.10

    _section_heading(document, "Projects")
    for item in profile.projects:
        paragraph = document.add_paragraph()
        paragraph.add_run(item.name).bold = True
        if item.url:
            paragraph.add_run(f" — {item.url}")
        document.add_paragraph(item.summary)
        technologies = document.add_paragraph()
        technologies.add_run("Technologies: ").bold = True
        technologies.add_run(", ".join(item.technologies))

    _section_heading(document, "Education")
    for item in profile.education:
        paragraph = document.add_paragraph()
        paragraph.add_run(item.qualification).bold = True
        paragraph.add_run(f", {item.institution} — {item.period}")

    document.save(destination)


def _section_heading(document: Document, title: str) -> None:
    document.add_paragraph(title, style="Heading 1")


def _configure_compact_reference_styles(document: Document) -> None:
    # Named override: portfolio_one_page_density. The compact-reference preset's
    # Letter geometry and style hierarchy remain, while body and section rhythm
    # are tightened so a concise fixture is not split into an orphan second page.
    normal = document.styles["Normal"]
    _set_style_font(normal, "Calibri", 10, RGBColor(32, 38, 48))
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.10

    heading_tokens = {
        "Heading 1": (13, HEADING_BLUE, 8, 4),
        "Heading 2": (13, HEADING_BLUE, 14, 7),
        "Heading 3": (12, HEADING_DARK_BLUE, 10, 5),
    }
    for style_name, (size, color, before, after) in heading_tokens.items():
        style = document.styles[style_name]
        _set_style_font(style, "Calibri", size, color, bold=True)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    list_style = document.styles["List Bullet"]
    _set_style_font(list_style, "Calibri", 10, RGBColor(32, 38, 48))
    list_style.paragraph_format.left_indent = Inches(0.375)
    list_style.paragraph_format.first_line_indent = Inches(-0.188)
    list_style.paragraph_format.space_after = Pt(2)
    list_style.paragraph_format.line_spacing = 1.10


def _set_style_font(style, name: str, size: int, color: RGBColor, *, bold: bool = False) -> None:
    style.font.name = name
    style._element.rPr.rFonts.set(qn("w:ascii"), name)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    style.font.size = Pt(size)
    style.font.color.rgb = color
    style.font.bold = bold


def _set_run_font(run, name: str = "Calibri") -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
